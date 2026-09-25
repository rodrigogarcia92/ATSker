"""The golden-set sheet, the operator's mapping, and one-shot validation (queue A.2).

`golden_run.py` and `golden_stats.py` share this because both read the sheet and
both resolve default paths against this script's own repository: a worktree of
`baseline-w2` has no `fixtures/` at all (it is gitignored), so the working
directory can never be the source of truth for the inputs. `src.*` is still read
from the working directory — the only import below is the user lookup, which
follows the same rule the runner enforces.

Validation reports every problem at once and never prints a mapping value: the
file holds one email and one filesystem path per row — personal data the
operator typed, not output. Problems name the sheet's own keys (`profile`,
`cv_version`, `pair_id`) and nothing else. (queue A.2, amended by
`docs/omp_queue_runner_resume.md` §2: in-scope rows only, blank `cv_version` is
one error listing its `pair_id`s, and mapping entries no in-scope row uses are
ignored entirely.)
"""

from __future__ import annotations

import subprocess
import tomllib
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.models import Profile, User

# The repository these scripts live in — where `fixtures/` is read from and
# written to regardless of which checkout is being measured (resume §3).
SCRIPT_REPO = Path(__file__).resolve().parents[1]
DEFAULT_SHEET = SCRIPT_REPO / "fixtures" / "golden" / "golden_set_project_26-2.xlsx"
DEFAULT_MAP = SCRIPT_REPO / "fixtures" / "golden" / "golden_run.local.toml"
DEFAULT_RUNS_DIR = SCRIPT_REPO / "fixtures" / "golden" / "runs"

# The columns the runner reads. The sheet has more (the rater's own export), and
# their order is the operator's; only these header names are the contract.
SHEET_COLUMNS = (
    "pair_id",
    "JD",
    "bucket",
    "score",
    "score_direct_only",
    "profile",
    "lang",
    "split",
    "rater",
    "cv_version",
)

# The CSV layout both scripts honour: one column per sub-score goes between the
# two blocks, read generically from the report. The stats reader identifies the
# sub-score columns as whatever is left between them.
#
# `vocabulary_hash` sits with the scorer's own version (F3.1): two runs are never
# compared across vocabularies without the report saying so. `model` and
# `reasoning_effort` are the provider configuration the calls were actually made
# with, read back from the report's `llm.calls` rather than from the flags.
COLUMNS_BEFORE_SUB_SCORES = (
    "pair_id",
    "profile",
    "split",
    "lang",
    "track",
    "weights_version",
    "vocabulary_hash",
    "score",
    "uncapped",
    "cap",
)
COLUMNS_AFTER_SUB_SCORES = (
    "skills_direct",
    "skills_transfer",
    "cost_usd",
    "cached_calls",
    "model",
    "reasoning_effort",
    "application_id",
    "commit",
)


class SheetError(Exception):
    """The sheet is missing, unreadable, or missing a column the runner needs."""


class MappingFileError(Exception):
    """The mapping file is missing or malformed. The message never carries its contents."""


@dataclass(frozen=True)
class Pair:
    """One (track, JD) pair as the sheet states it. `score` is the rater's.

    `rater` names who gave it, so a report can say whose judgement it is agreeing
    with rather than hard-coding one name (A.6).
    """

    pair_id: str
    jd: str
    bucket: str
    score: float | None
    score_direct_only: float | None
    profile: str
    lang: str
    split: str
    cv_version: str
    rater: str = ""


@dataclass(frozen=True)
class Mapping:
    """The operator's two lookups: sheet key -> email, sheet key -> CV file."""

    profiles: dict[str, str] = field(default_factory=dict)
    cvs: dict[str, str] = field(default_factory=dict)
    source: Path = DEFAULT_MAP


@dataclass(frozen=True)
class InputCheck:
    """Every problem found at once, plus what did resolve (never the mapped values)."""

    problems: tuple[str, ...]
    users: dict[str, User] = field(default_factory=dict)
    profiles: dict[str, Profile] = field(default_factory=dict)
    files: dict[str, Path] = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not self.problems


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def output_is_ignored(path: Path) -> bool:
    """Queue A.3 step 5: run output may only land where `git check-ignore` matches.

    The guard runs against the repository the file actually lands in (resume §3),
    which is why it starts in the output's own parent folder.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        ["git", "-C", str(path.parent), "check-ignore", "-q", str(path)],
        check=False,
        capture_output=True,
    )
    return result.returncode == 0


def _number(value: object) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    return float(value)


def _cell(row: tuple[object, ...], columns: dict[str, int], name: str) -> object:
    index = columns[name]
    return row[index] if index < len(row) else None


def read_pairs(path: Path, *, include_test: bool) -> list[Pair]:
    """The sheet's in-scope rows: `dev`, plus `test` only with `include_test`.

    Columns are found by header name (resume §1). Rows whose split is anything
    else are out of scope and never read — the runner never writes to the sheet.
    """
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError as exc:
        raise SheetError(f"sheet not found: {path}") from exc
    except Exception as exc:  # noqa: BLE001 - any reader failure is one message to the operator
        raise SheetError(f"sheet cannot be read: {path}") from exc
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise SheetError(f"sheet has no header row: {path}")
        columns = {_text(name): index for index, name in enumerate(header) if name is not None}
        missing = [name for name in SHEET_COLUMNS if name not in columns]
        if missing:
            raise SheetError("sheet is missing columns: " + ", ".join(missing))
        pairs: list[Pair] = []
        for row in rows:
            if all(cell is None for cell in row):
                continue
            split = _text(_cell(row, columns, "split")).lower()
            if split != "dev" and not (include_test and split == "test"):
                continue
            pairs.append(
                Pair(
                    pair_id=_text(_cell(row, columns, "pair_id")),
                    jd=_text(_cell(row, columns, "JD")),
                    bucket=_text(_cell(row, columns, "bucket")),
                    score=_number(_cell(row, columns, "score")),
                    score_direct_only=_number(_cell(row, columns, "score_direct_only")),
                    profile=_text(_cell(row, columns, "profile")),
                    lang=_text(_cell(row, columns, "lang")),
                    split=split,
                    cv_version=_text(_cell(row, columns, "cv_version")),
                    rater=_text(_cell(row, columns, "rater")),
                )
            )
        return pairs
    finally:
        workbook.close()


def _table(data: dict[str, object], name: str, source: Path) -> dict[str, str]:
    section = data.get(name)
    if section is None:
        return {}
    if not isinstance(section, dict) or any(
        not isinstance(key, str) or not isinstance(value, str) for key, value in section.items()
    ):
        # The message names the section, never a value from the file.
        raise MappingFileError(f"{source}: [{name}] must map names to strings")
    return dict(section)


def load_mapping(path: Path) -> Mapping:
    """The two tables the runner needs. A missing file is an error, not an empty map."""
    try:
        raw = path.read_bytes()
    except FileNotFoundError as exc:
        raise MappingFileError(f"mapping file not found: {path}") from exc
    try:
        data = tomllib.loads(raw.decode("utf-8"))
    except (tomllib.TOMLDecodeError, UnicodeDecodeError) as exc:
        # TOMLDecodeError quotes the source line; the operator sees only this.
        raise MappingFileError(f"mapping file is not valid TOML: {path}") from exc
    return Mapping(
        profiles=_table(data, "profiles", path),
        cvs=_table(data, "cvs", path),
        source=path,
    )


def check_inputs(pairs: list[Pair], mapping: Mapping, session: Session) -> InputCheck:
    """Every problem at once (queue A.2): unmapped keys, missing files, unknown users.

    Only what in-scope rows use is checked (resume §2) — an entry nothing needs
    is never looked up against the database and its file is never touched.
    """
    problems: list[str] = []
    users: dict[str, User] = {}
    profiles: dict[str, Profile] = {}
    files: dict[str, Path] = {}

    blank = [pair.pair_id for pair in pairs if not pair.cv_version]
    if blank:
        problems.append("blank cv_version in rows: " + ", ".join(blank))

    for name in dict.fromkeys(pair.profile for pair in pairs):
        email = mapping.profiles.get(name)
        if email is None:
            problems.append(f"profile {name!r} has no entry in [profiles]")
            continue
        user = session.scalar(select(User).where(User.email == email))
        if user is None:
            # The email is mapping content: named by what it fails, never printed.
            problems.append(f"profile {name!r}: no user with the mapped email")
            continue
        profile = session.scalar(select(Profile).where(Profile.user_id == user.id))
        if profile is None:
            problems.append(f"profile {name!r}: the mapped user has no profile row")
            continue
        users[name] = user
        profiles[name] = profile

    for version in dict.fromkeys(pair.cv_version for pair in pairs if pair.cv_version):
        raw = mapping.cvs.get(version)
        if raw is None:
            problems.append(f"cv_version {version!r} has no entry in [cvs]")
            continue
        candidate = Path(raw).expanduser()
        if not candidate.is_absolute():
            candidate = mapping.source.parent / candidate
        if not candidate.is_file():
            problems.append(f"cv_version {version!r}: the mapped CV file does not exist")
        else:
            files[version] = candidate

    return InputCheck(problems=tuple(problems), users=users, profiles=profiles, files=files)
