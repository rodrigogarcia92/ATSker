"""Synthetic fixtures shared by the golden-set runner's tests.

Everything here is synthetic — sheets, mappings, CV bytes and canned provider
answers. The real sheet and the operator's mapping hold personal data and no
test reads them. The provider is stubbed at `src.extract.client._post`, the same
seam the rest of the suite uses, so the whole run happens offline.
"""

import json
import subprocess
import sys
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import golden_inputs  # noqa: E402

# One extraction answer for any run: one role, one bullet, two credentials.
RESUME_PROPOSAL = json.dumps(
    {
        "language": "es",
        "headline": "Analista de datos",
        "summary": "Cinco años automatizando reportes.",
        "roles": [
            {
                "employer": "Globex",
                "title": "Data Scientist",
                "start_date": "2021-01",
                "bullets": ["Modelé el riesgo de crédito con Python"],
            }
        ],
        "credentials": [
            {
                "kind": "education",
                "name": "Bachiller en Ingeniería de Sistemas",
                "issuer": "Universidad Nacional de Ingeniería",
                "end_date": "2019",
            },
            {
                "kind": "certification",
                "name": "Scrum Master",
                "issuer": "Scrum Alliance",
                "issued_date": "2020-03",
            },
        ],
        "skills": ["python", "cobalto"],
    }
)

# One posting shape for every analysed pair: two requirements, in this order.
JD_PROPOSAL = json.dumps(
    {
        "language": "es",
        "title": "Analista de datos",
        "seniority": "mid",
        "years_required": 3,
        "must_have": ["python"],
        "nice_to_have": ["machine learning"],
        "raw_keywords": ["reportes"],
        "hard_requirements": ["título universitario"],
        "location_policy": "hybrid_lima",
    }
)

QUERIES = json.dumps(
    {
        "queries": [
            {"label": "broad", "expr": {"and": ["analista de datos", "python"]}},
            {
                "label": "strict",
                "expr": {"and": ["python", "machine learning", {"not": ["excel"]}]}},
        ]
    }
)

# Both requirements graded level 1 from the summary — the non-bullet source that
# needs no bullet id, so one answer fits a run whose bullets do not exist yet.
GRADING = json.dumps(
    {
        "grades": [
            {
                "index": 0,
                "level": 1,
                "basis": "direct",
                "source": "summary",
                "evidence_id": None,
                "transfer_reason": None,
            },
            {
                "index": 1,
                "level": 1,
                "basis": "direct",
                "source": "summary",
                "evidence_id": None,
                "transfer_reason": None,
            },
        ],
        "note": "una nota",
    }
)

# Exactly the gaps the grading above leaves, in detection's band order: the
# must-have first, then the nice-to-have.
GAP_ANSWER = json.dumps(
    {
        "gaps": [
            {
                "requirement": "python",
                "blocking": True,
                "interpretation": "Python aparece en el resumen.",
            },
            {
                "requirement": "machine learning",
                "blocking": False,
                "interpretation": "Machine learning está mencionado sin una historia.",
            },
        ],
        "warning": {
            "text": "Salto hacia banca privada",
            "handling": "Llévalo como transferencia de método analítico.",
        },
        "not_found": "No hay historias sobre despliegue de modelos.",
    }
)


class Provider:
    """Every call of one run, by schema name. Unknown names get the résumé answer."""

    def __init__(self) -> None:
        self.calls: list[dict] = []
        self.answers = {
            "JDProposal": JD_PROPOSAL,
            "BooleanQuerySet": QUERIES,
            "EvidenceGrading": GRADING,
            "GapInterpretation": GAP_ANSWER,
        }

    def __call__(self, payload: dict, api_key: str) -> dict:
        self.calls.append(payload)
        name = payload["response_format"]["json_schema"]["name"]
        return {
            "choices": [{"message": {"content": self.answers.get(name, RESUME_PROPOSAL)}}],
            "usage": {"prompt_tokens": 2200, "completion_tokens": 140, "cost": "0.000123"},
        }


def write_sheet(
    path: Path,
    rows: list[dict[str, object]],
    columns: tuple[str, ...] = golden_inputs.SHEET_COLUMNS,
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(columns))
    for entry in rows:
        sheet.append([entry.get(name) for name in columns])
    workbook.save(path)
    return path


def write_map(path: Path, profiles: dict[str, str], cvs: dict[str, str]) -> Path:
    lines = ["[profiles]"]
    lines += [f"{json.dumps(name)} = {json.dumps(email)}" for name, email in profiles.items()]
    lines += ["", "[cvs]"]
    lines += [f"{json.dumps(name)} = {json.dumps(value)}" for name, value in cvs.items()]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def row(
    pair_id: object,
    *,
    profile: str = "Ana",
    cv_version: str = "CV-1",
    split: str = "dev",
    rater: str = "rater_v2",
    **extra: object,
) -> dict[str, object]:
    return {
        "pair_id": pair_id,
        "JD": f"posting {pair_id}",
        "bucket": "plausible",
        "score": 3.5,
        "score_direct_only": 2.5,
        "profile": profile,
        "lang": "es",
        "split": split,
        "rater": rater,
        "cv_version": cv_version,
        **extra,
    }


def docx_bytes(*paragraphs: str) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def gitignored_dir(tmp_path: Path) -> Path:
    """A throwaway repository whose paths are all gitignored, for the output guard."""
    repo = tmp_path / "repo"
    repo.mkdir()
    subprocess.run(["git", "init"], cwd=repo, check=True, capture_output=True)
    (repo / ".gitignore").write_text("*\n", encoding="utf-8")
    return repo


@contextmanager
def session_from(session_factory):
    """The runner's `session_scope` seam: one test-database session per run."""
    with session_factory() as db:
        yield db
