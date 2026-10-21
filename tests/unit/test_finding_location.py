"""F5.1.1's "show me where": the location view, its characters, and its limits.

The point of these tests is the promise decision 9 makes: the stored finding carries
boxes and no text, and the characters of a hidden word are read from the archived
original at request time, inside the owner's session, and stored nowhere. So the same
word that the page view renders must be absent from the résumé's gate, from the stored
report, and from the log.
"""

import json
import logging
import re
from decimal import Decimal
from io import BytesIO

from docx import Document
from pdfplumber.utils.exceptions import MalformedPDFException
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from src.applications import create_application
from src.extract import client as extract_client
from src.models import Resume
from src.web.routes import ingest as ingest_routes
from tests.conftest import login

HIDDEN = "confidencial9"
PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"

PROPOSAL = json.dumps(
    {
        "language": "es",
        "headline": "Analista de datos",
        "summary": "Cinco años automatizando reportes.",
        "roles": [
            {
                "employer": "Globex",
                "title": "Data Scientist",
                "start_date": "2021-01",
                "bullets": ["Modelé el riesgo de crédito con Python"],
            }
        ],
        "credentials": [],
        "skills": ["python"],
    }
)


def provider(monkeypatch):
    calls: list[dict] = []

    def stub(payload, api_key):
        calls.append(payload)
        return {
            "choices": [{"message": {"content": PROPOSAL}}],
            "usage": {"prompt_tokens": 1500, "completion_tokens": 200, "cost": "0.000090"},
        }

    monkeypatch.setattr(extract_client, "_post", stub)
    return calls


def white_text_pdf(text: str = HIDDEN) -> bytes:
    """A readable page carrying one word in white — an F5.1.2 blocker."""
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    pdf.setFont("Helvetica", 10)
    pdf.setFillColorRGB(0, 0, 0)
    for index in range(6):
        pdf.drawString(60, 720 - 20 * index, f"Experience {index}: Python and SQL reporting")
    pdf.setFillColorRGB(1, 1, 1)
    pdf.drawString(60, 560, text)
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def email_header_docx() -> bytes:
    document = Document()
    document.add_paragraph("Experiencia: analista de datos con Python")
    document.sections[0].header.paragraphs[0].text = "persona2@example.com"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload(client, data: bytes, name: str, track_id: str | None = None):
    content_type = (
        "application/pdf"
        if name.endswith(".pdf")
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    return client.post(
        "/import",
        files={"file": (name, data, content_type)},
        data={"track_id": track_id} if track_id else None,
        follow_redirects=False,
    )


def finding_index(resume: Resume, check: str) -> int:
    return next(
        index
        for index, finding in enumerate(resume.gate["findings"])
        if finding["check"] == check
    )


def plain(html: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html))


def test_the_location_view_renders_the_page_and_the_hidden_characters(
    client, session, career, monkeypatch, caplog
):
    provider(monkeypatch)
    login(client, career.first.email)
    # INFO, not DEBUG: pdfminer logs the decoded content stream at DEBUG, and that is
    # the library's log, not the application's. The promise is that ATSker logs no
    # document text (PRD 10) — so the level the app runs at is the one asserted.
    caplog.set_level(logging.INFO)

    uploaded = upload(client, white_text_pdf(), "cv.pdf", track_id=str(career.first_track.id))
    assert uploaded.status_code == 303
    resume = session.query(Resume).one()
    index = finding_index(resume, "invisible_text")

    review = client.get(f"/import/{resume.id}").text
    link = f"/import/{resume.id}/findings/{index}"
    assert link in review, "a word-level finding offers its location (F5.1.1)"
    assert "Mostrarme dónde" in review

    view = client.get(link)
    assert view.status_code == 200
    assert HIDDEN in plain(view.text), "the page view reads the hidden word at request time"

    image = client.get(f"{link}/page.png")
    assert image.status_code == 200
    assert image.headers["content-type"] == "image/png"
    assert image.headers["cache-control"] == "no-store"
    assert image.content.startswith(PNG_SIGNATURE)

    assert HIDDEN not in json.dumps(resume.gate), "the stored finding keeps boxes, no text"
    assert HIDDEN not in caplog.text, "no résumé text in a log"

    # A page PDFium refuses is a 404: the image hides itself and the text says so.
    image = re.search(r'<img[^>]*src="[^"]*/page\.png"[^>]*>\s*(<p[^>]*>[^<]*</p>)', view.text)
    assert image is not None
    assert 'onerror="this.hidden=true;this.nextElementSibling.hidden=false"' in image.group(0)
    assert re.search(r"·\s*1</p>", image.group(1)), "the fallback names the page"


def test_another_profile_cannot_open_the_location_view_or_its_image(
    client, session, career, monkeypatch
):
    provider(monkeypatch)
    login(client, career.first.email)
    uploaded = upload(client, white_text_pdf(), "cv.pdf")
    assert uploaded.status_code == 303
    resume = session.query(Resume).one()
    index = finding_index(resume, "invisible_text")
    client.post("/logout")

    login(client, career.second.email)
    assert client.get(f"/import/{resume.id}/findings/{index}").status_code == 404
    assert client.get(f"/import/{resume.id}/findings/{index}/page.png").status_code == 404


def test_the_location_view_is_never_cached(client, session, career, monkeypatch):
    """N6.6: the HTML view can carry the hidden characters, so it is as uncacheable as
    the image it embeds."""
    provider(monkeypatch)
    login(client, career.first.email)
    upload(client, white_text_pdf(), "cv.pdf")
    resume = session.query(Resume).one()
    index = finding_index(resume, "invisible_text")

    view = client.get(f"/import/{resume.id}/findings/{index}")

    assert view.status_code == 200
    assert view.headers["cache-control"] == "no-store"


def test_a_page_the_file_does_not_have_is_a_404(client, session, career, monkeypatch):
    """N6.6: a stored page number is not proof the page exists."""
    provider(monkeypatch)
    login(client, career.first.email)
    upload(client, white_text_pdf(), "cv.pdf")
    resume = session.query(Resume).one()
    index = finding_index(resume, "invisible_text")

    findings = [dict(finding) for finding in resume.gate["findings"]]
    findings[index] = {
        **findings[index],
        "detail": {**findings[index]["detail"], "page": 9},
    }
    resume.gate = {**resume.gate, "findings": findings}
    session.commit()

    link = f"/import/{resume.id}/findings/{index}"
    assert client.get(f"{link}/page.png").status_code == 404
    assert client.get(link).status_code == 404


def test_a_check_with_no_box_geometry_has_no_page_image(client, session, career, monkeypatch):
    """N6.6: the PNG route checks that the check is locatable, not only the index."""
    provider(monkeypatch)
    login(client, career.first.email)
    upload(client, white_text_pdf(), "cv.pdf")
    resume = session.query(Resume).one()
    index = finding_index(resume, "invisible_text")

    findings = [dict(finding) for finding in resume.gate["findings"]]
    findings[index] = {**findings[index], "check": "multi_column"}
    resume.gate = {**resume.gate, "findings": findings}
    session.commit()

    link = f"/import/{resume.id}/findings/{index}"
    assert client.get(f"{link}/page.png").status_code == 404
    assert client.get(link).status_code == 404


def test_a_renderer_that_fails_here_is_a_404_and_one_warning(
    client, session, career, monkeypatch, caplog
):
    """The pane's serialised PDFium can still refuse a file: no image, never a 500."""
    provider(monkeypatch)
    login(client, career.first.email)
    upload(client, white_text_pdf(), "cv.pdf")
    resume = session.query(Resume).one()
    index = finding_index(resume, "invisible_text")

    def failing(*_args, **_kwargs) -> bytes:
        raise MalformedPDFException("Failed to load document (PDFium: Data format error).")

    monkeypatch.setattr(ingest_routes, "render_page_png", failing)
    caplog.set_level(logging.WARNING)

    answer = client.get(f"/import/{resume.id}/findings/{index}/page.png")

    assert answer.status_code == 404
    assert f"resume={resume.id}" in caplog.text


def test_a_docx_finding_says_the_location_view_is_pdf_only(
    client, session, career, monkeypatch
):
    provider(monkeypatch)
    login(client, career.first.email)

    uploaded = upload(client, email_header_docx(), "cv.docx")
    assert uploaded.status_code == 303
    resume = session.query(Resume).one()
    assert finding_index(resume, "contact_in_margin") >= 0

    review = client.get(f"/import/{resume.id}").text
    assert "solo está disponible para PDF" in review
    assert "/findings/" not in review, "a DOCX finding offers no location link"


def test_the_stored_report_links_through_the_source_resume_without_the_characters(
    client, session, career, monkeypatch
):
    provider(monkeypatch)
    login(client, career.first.email)

    uploaded = upload(client, white_text_pdf(), "cv.pdf", track_id=str(career.first_track.id))
    assert uploaded.status_code == 303
    resume = session.query(Resume).one()
    index = finding_index(resume, "invisible_text")

    application = create_application(
        career.first_profile.id,
        session,
        jd_text="Oferta de ejemplo",
        jd_hash="sha256:n6-e3",
        report=stored_report(resume.gate),
        weights_version="w4",
        prompt_versions={"jd_extract": "jd_extract_v2"},
        token_cost_usd=Decimal("0"),
        track_id=career.first_track.id,
        scan_readability=resume.gate,
    )
    session.commit()

    report_page = client.get(f"/applications/{application.id}").text
    link = f"/import/{resume.id}/findings/{index}"
    assert link in report_page, "the report links through the track's source résumé"
    assert HIDDEN not in json.dumps(application.report)

    assert HIDDEN in plain(client.get(link).text)


def stored_report(gate: dict) -> dict:
    """The smallest report the page renders: enough keys, nothing invented."""
    return {
        "jd": {
            "language": "es",
            "title": "Analista de datos",
            "seniority": "mid",
            "years_required": 3,
            "must_have": ["python"],
            "nice_to_have": [],
            "raw_keywords": ["reportes"],
            "hard_requirements": [],
            "location_policy": "hybrid_lima",
        },
        "grades": [],
        "grading_note": None,
        "gate": gate,
        "recall": {"checked": False, "queries": []},
        "score": {
            "weights_version": "w4",
            "weights_hash": "sha256:0",
            "weighted": 0.4,
            "uncapped": 40,
            "capped": 40,
            "cap": 40,
            "cap_reason": "invisible_text",
            "unavailable": {},
            "effective_weights": {
                "skills": 0.4,
                "experience": 0.25,
                "title": 0.2,
                "keywords": 0.15,
            },
            "effective_weights_hash": "sha256:0",
            "sub_scores": {},
        },
        "llm": {"cost_usd": "0", "calls": []},
    }
