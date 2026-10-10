"""The import flow over HTTP (PRD F2.1-F2.5): upload, review, diff, merge.

The provider is stubbed at `client._post`, so the whole front door runs offline:
gate findings, the editable draft, the diff screen, and the merge that is the only
thing allowed to touch career data.
"""

import json
import re
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import NullPool

from src.auth.accounts import create_account
from src.bullets import list_bullets
from src.certifications import list_certifications
from src.education import list_educations
from src.extract import client as extract_client
from src.extract.client import TransactionHeldAcrossProviderCall
from src.models import Application, Base, LlmCall, Resume, Track
from src.roles import list_roles
from src.tracks import track_bullets
from src.web.deps import get_db
from src.web.main import create_app
from tests.conftest import PASSWORD, login

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

POSTING = """
Analista de datos — Banco Demo 1 (híbrido, Lima)
Buscamos un analista con 3 años de experiencia en Python y machine learning.
Imprescindible: título universitario.
"""

PROPOSAL = json.dumps(
    {
        "language": "es",
        "headline": "Analista de datos",
        "summary": "Cinco años automatizando reportes.",
        "roles": [
            {
                "employer": "Globex",
                "title": "Data Scientist",
                "start_date": "2021-01",
                "bullets": ["Modelé el riesgo de crédito con Python"],
            }
        ],
        "credentials": [
            {
                "kind": "education",
                "name": "Bachiller en Ingeniería de Sistemas",
                "issuer": "Universidad Nacional de Ingeniería",
                "end_date": "2019",
            },
            {
                "kind": "certification",
                "name": "Scrum Master",
                "issuer": "Scrum Alliance",
                "issued_date": "2020-03",
            },
        ],
        "skills": ["python", "cobalto"],
    }
)

REVIEWED_FORM = {
    "language": "es",
    "headline": "Analista de datos",
    "summary": "Cinco años automatizando reportes.",
    "skills": "python",
    "role_count": "1",
    "role_0_employer": "Globex",
    "role_0_title": "Data Scientist",
    "role_0_start": "2021-01-01",
    "role_0_end": "",
    "bullet_count_0": "1",
    "bullet_0_0_text": "Modelé el riesgo de crédito con Python",
    "bullet_0_0_chips": "python",
    "credential_count": "2",
    "credential_0_kind": "education",
    "credential_0_name": "Bachiller en Ingeniería de Sistemas",
    "credential_0_issuer": "Universidad Nacional de Ingeniería",
    "credential_0_field": "",
    "credential_0_start": "",
    "credential_0_end": "2019-01-01",
    "credential_0_issued": "",
    "credential_0_expiry": "",
    "credential_0_credential_id": "",
    "credential_0_credential_url": "",
    "credential_0_review": "off",
    "credential_1_kind": "certification",
    "credential_1_name": "Scrum Master",
    "credential_1_issuer": "Scrum Alliance",
    "credential_1_field": "",
    "credential_1_start": "",
    "credential_1_end": "",
    "credential_1_issued": "2020-03-01",
    "credential_1_expiry": "",
    "credential_1_credential_id": "",
    "credential_1_credential_url": "",
    "credential_1_review": "off",
}

MERGE_FORM = {
    "merge_role_0": "on",
    "merge_bullet_0_0": "on",
    "merge_education_0": "on",
    "merge_certification_0": "on",
    "merge_skill_0": "on",
}


class Provider:
    """One résumé extraction answer, priced like a real call."""

    def __init__(self, payload: str = PROPOSAL):
        self.calls: list[dict] = []
        self.payload = payload

    def __call__(self, payload: dict, api_key: str) -> dict:
        self.calls.append(payload)
        return {
            "choices": [{"message": {"content": self.payload}}],
            "usage": {"prompt_tokens": 1500, "completion_tokens": 200, "cost": "0.000090"},
        }


@pytest.fixture
def provider(monkeypatch, test_settings) -> Provider:
    stub = Provider()
    monkeypatch.setattr(extract_client, "_post", stub)
    return stub


def docx_bytes(*paragraphs: str, header: str | None = None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if header is not None:
        document.sections[0].header.paragraphs[0].text = header
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pdf_bytes(*lines: str) -> bytes:
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    pdf.setFont("Helvetica", 10)
    y = 720
    for text in lines:
        pdf.drawString(60, y, text)
        y -= 20
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def plain(html: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html))


def upload(
    client,
    data: bytes,
    name: str = "cv.docx",
    content_type: str = DOCX_TYPE,
    track_id: str | None = None,
):
    return client.post(
        "/import",
        files={"file": (name, data, content_type)},
        data={"track_id": track_id} if track_id else None,
        follow_redirects=False,
    )


def test_upload_review_merge_round_trip(client, session, career, provider):
    login(client, career.first.email)
    before = len(list_roles(career.first_profile.id, session, active_only=False))

    uploaded = upload(client, docx_bytes("Experiencia", "Reduje el cierre un 30% con Python"))
    assert uploaded.status_code == 303
    review_url = uploaded.headers["location"]

    page = plain(client.get(review_url).text)
    assert "Lectura del documento" in page, "findings come before any extracted content"
    assert "Modelé el riesgo de crédito con Python" in page, "the draft is editable"
    assert session.query(Resume).count() == 1
    assert len(list_roles(career.first_profile.id, session, active_only=False)) == before

    merge_page = plain(client.post(f"{review_url}/review", data=REVIEWED_FORM).text)
    assert "Nuevo" in merge_page
    assert "Modelé el riesgo de crédito con Python" in merge_page

    merged = plain(client.post(f"{review_url}/merge", data=MERGE_FORM).text)
    assert "Importación aplicada" in merged

    profile_id = career.first_profile.id
    roles = [
        role
        for role in list_roles(profile_id, session, active_only=False)
        if role.employer == "Globex"
    ]
    assert [role.title for role in roles] == ["Data Scientist"]
    bullets = list_bullets(profile_id, session, active_only=True, role_id=roles[0].id)
    assert [bullet.text_es for bullet in bullets] == ["Modelé el riesgo de crédito con Python"]
    assert [entry.name for entry in list_certifications(profile_id, session)] == ["Scrum Master"]
    assert [entry.institution for entry in list_educations(profile_id, session)] == [
        "Universidad Nacional de Ingeniería"
    ]
    session.expire_all()
    assert career.first_profile.skills == ["python"]
    assert session.query(Resume).one().merged_at is not None


def test_the_same_document_uploaded_twice_is_cached_and_diffs_as_unchanged(
    client, session, career, provider
):
    login(client, career.first.email)
    data = docx_bytes("Experiencia", "Reduje el cierre un 30% con Python")

    first = upload(client, data)
    client.post(f"{first.headers['location']}/review", data=REVIEWED_FORM)
    client.post(f"{first.headers['location']}/merge", data=MERGE_FORM)

    second = upload(client, data)
    assert len(provider.calls) == 1, "the extraction is cached by content hash"

    merge_page = plain(client.post(f"{second.headers['location']}/review", data=REVIEWED_FORM).text)
    assert "0 nuevos" in merge_page
    assert "0 cambian" in merge_page

    merged = client.post(f"{second.headers['location']}/merge", data=MERGE_FORM)
    assert merged.status_code == 200

    profile_id = career.first_profile.id
    globex = [
        role
        for role in list_roles(profile_id, session, active_only=False)
        if role.employer == "Globex"
    ]
    assert len(globex) == 1
    assert len(list_bullets(profile_id, session, active_only=False)) == 2, "nothing duplicated"
    assert len(list_certifications(profile_id, session)) == 1
    assert len(list_educations(profile_id, session)) == 1


def test_the_review_screen_can_move_an_entry_between_education_and_certification(
    client, session, career, provider
):
    """The classifier's rules can fire wrong; the kind is a field the user can fix."""
    login(client, career.first.email)
    uploaded = upload(client, docx_bytes("Experiencia", "Reduje el cierre un 30% con Python"))
    review_url = uploaded.headers["location"]

    fixed = dict(REVIEWED_FORM, **{"credential_0_kind": "certification"})
    merge_page = plain(client.post(f"{review_url}/review", data=fixed).text)
    assert "0 nuevos" not in merge_page

    merged = client.post(
        f"{review_url}/merge",
        data=dict(MERGE_FORM, **{"merge_certification_1": "on"}),
    )
    assert merged.status_code == 200

    profile_id = career.first_profile.id
    assert list_educations(profile_id, session) == []
    names = [entry.name for entry in list_certifications(profile_id, session)]
    assert sorted(names) == ["Bachiller en Ingeniería de Sistemas", "Scrum Master"]


def test_the_review_screen_shows_findings_before_the_draft(client, session, career, provider):
    login(client, career.first.email)

    uploaded = upload(
        client,
        docx_bytes("Experiencia", header="persona2@example.com | +00 111 222 333"),
    )

    page = plain(client.get(uploaded.headers["location"]).text)
    assert "Bloqueante" in page
    assert "persona2@example.com" not in page, "findings never echo the text they matched"


def test_an_overlong_headline_warns_on_review_and_is_not_set_by_the_merge(
    client, session, career, monkeypatch
):
    """The review screen warns under the field so the user can shorten the headline
    before merging; when it is merged anyway, the merged page repeats the note and
    the track headline stays unset.
    """
    long_headline = "x" * 210
    stub = Provider(json.dumps({**json.loads(PROPOSAL), "headline": long_headline}))
    monkeypatch.setattr(extract_client, "_post", stub)
    login(client, career.first.email)
    uploaded = upload(client, docx_bytes("Experiencia", "Reduje el cierre un 30% con Python"))
    review_url = uploaded.headers["location"]

    review = plain(client.get(review_url).text)
    assert "Titular no establecido: 210 caracteres, límite 200" in review

    client.post(f"{review_url}/review", data=dict(REVIEWED_FORM, headline=long_headline))
    merged = plain(client.post(f"{review_url}/merge", data=MERGE_FORM).text)

    assert "Titular no establecido: 210 caracteres, límite 200" in merged
    track = session.scalar(select(Track).where(Track.name == "cv"))
    assert track.headline_es is None, "headline stays unset"


def test_upload_without_a_model_keeps_the_findings_and_stores_nothing(
    client, session, career, monkeypatch
):
    from src.config import get_settings

    # An empty value, not a deleted one: a developer's .env would otherwise supply
    # a real model and the test would leave for the network.
    monkeypatch.setenv("OPENROUTER_MODEL", "")
    get_settings.cache_clear()
    try:
        login(client, career.first.email)
        response = upload(client, docx_bytes("Experiencia", "Reduje el cierre un 30%"))
    finally:
        get_settings.cache_clear()

    assert response.status_code == 200
    assert "OPENROUTER_MODEL" in plain(response.text)
    assert "Lectura del documento" in plain(response.text)
    assert session.query(Resume).count() == 0


def test_a_broken_docx_is_refused_by_the_reader(client, session, career, provider):
    login(client, career.first.email)

    response = upload(client, b"this is not a zip file", name="cv.docx")

    assert response.status_code == 200
    assert "no se pudo leer" in plain(response.text)
    assert session.query(Resume).count() == 0
    assert provider.calls == []


def test_other_formats_are_refused_by_name(client, session, career, provider):
    login(client, career.first.email)

    response = upload(client, b"plain text", name="cv.txt")

    assert response.status_code == 200
    assert "Solo se aceptan PDF o DOCX" in plain(response.text)
    assert session.query(Resume).count() == 0


def test_another_profiles_import_is_404(client, session, career, provider):
    login(client, career.first.email)
    uploaded = upload(client, docx_bytes("Experiencia", "Reduje el cierre un 30% con Python"))
    review_url = uploaded.headers["location"]

    login(client, career.second.email)

    assert client.get(review_url).status_code == 404
    assert client.post(f"{review_url}/review", data=REVIEWED_FORM).status_code == 404
    assert client.post(f"{review_url}/merge", data=MERGE_FORM).status_code == 404
    assert client.post(f"{review_url}/discard").status_code == 404
    resume = session.query(Resume).one()
    assert resume.merged_at is None


def test_a_provider_call_slower_than_the_transaction_keeps_the_upload_and_the_row(
    tmp_path, test_settings, monkeypatch
):
    """Neon kills a connection left idle in a transaction while the provider works.

    The upload must still succeed and the paid call must still reach the ledger: the
    call site releases its reads before the network call, and the row is written on
    the client's own connection afterwards.
    """
    engine = create_engine(f"sqlite:///{tmp_path / 'slow-provider.db'}", poolclass=NullPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as setup:
        create_account(setup, "persona2@example.com", PASSWORD, "Ana")
        setup.commit()

    opened: list[object] = []

    @event.listens_for(engine, "checkout")
    def _remember(dbapi_connection, connection_record, connection_proxy):
        opened.append(dbapi_connection)

    def slow_provider(payload: dict, api_key: str) -> dict:
        # The provider is slow, and the database kills every connection the request
        # has used so far. The ledger write happens after this, on a connection of
        # its own: the paid answer must survive a dead request connection.
        for connection in opened:
            connection.close()
        return {
            "choices": [{"message": {"content": PROPOSAL}}],
            "usage": {"prompt_tokens": 1500, "completion_tokens": 200, "cost": "0.000090"},
        }

    app = create_app()

    def override_get_db():
        with factory() as db:
            try:
                yield db
            except Exception:
                db.rollback()
                raise

    app.dependency_overrides[get_db] = override_get_db
    monkeypatch.setattr(extract_client, "_post", slow_provider)

    with TestClient(app) as slow_client:
        login(slow_client, "persona2@example.com")
        uploaded = upload(
            slow_client, docx_bytes("Experiencia", "Reduje el cierre un 30% con Python")
        )

    assert uploaded.status_code == 303, "a slow provider call must not 500 the upload"
    with factory() as check:
        assert check.query(LlmCall).count() == 1, "the paid call keeps its ledger row"
        assert check.query(Resume).count() == 1


def test_a_route_that_never_releases_its_reads_is_refused(
    session_factory, career, test_settings, monkeypatch
):
    """The guard is the enforcement: with the release removed, the request stops
    before the provider is reached instead of holding a transaction across it."""
    app = create_app()

    def forgetful_get_db():
        with session_factory() as db:
            db.commit = lambda: None  # a route that never releases its reads
            try:
                yield db
            except Exception:
                db.rollback()
                raise

    app.dependency_overrides[get_db] = forgetful_get_db

    def forbidden(payload: dict, api_key: str) -> dict:
        pytest.fail("the provider must not be reached while a transaction is held")

    monkeypatch.setattr(extract_client, "_post", forbidden)

    with TestClient(app) as forgetful:
        login(forgetful, career.first.email)
        with pytest.raises(TransactionHeldAcrossProviderCall, match="resume_extract"):
            upload(forgetful, docx_bytes("Experiencia", "Reduje el cierre un 30% con Python"))


def test_a_docx_review_names_the_checks_that_did_not_run(client, session, career, provider):
    """A DOCX is weaker evidence: the panel says which checks it cannot answer."""
    login(client, career.first.email)

    uploaded = upload(client, docx_bytes("Experiencia", "Reduje el cierre un 30% con Python"))
    page = plain(client.get(uploaded.headers["location"]).text)

    assert "No se ejecutaron para este formato" in page
    assert "Bullets como imagen" in page
    assert "Fuentes no incrustadas" in page


def test_a_pdf_review_lists_no_skipped_checks(client, session, career, provider):
    login(client, career.first.email)

    uploaded = upload(
        client,
        pdf_bytes("Experiencia", "Reduje el cierre un 30% con Python"),
        name="cv.pdf",
        content_type="application/pdf",
    )
    page = plain(client.get(uploaded.headers["location"]).text)

    assert "No se ejecutaron para este formato" not in page


def test_an_unrecorded_call_keeps_the_draft_and_shows_the_warning(
    tmp_path, client, session, career, provider, monkeypatch
):
    """The paid answer is kept; the missing ledger row is said out loud, not hidden."""
    good = extract_client._client_session
    broken = create_engine(f"sqlite:///{tmp_path}")  # a directory: nothing can connect
    opened: list[int] = []

    def flaky(bind_session):
        opened.append(1)
        return good(bind_session) if len(opened) == 1 else Session(bind=broken)

    monkeypatch.setattr(extract_client, "_client_session", flaky)

    login(client, career.first.email)
    response = upload(client, docx_bytes("Experiencia", "Reduje el cierre un 30% con Python"))

    assert response.status_code == 200, "the answer is kept, so this is a warning, not an error"
    page = plain(response.text)
    assert "no se pudo registrar" in page
    assert "Modelé el riesgo de crédito con Python" in page, "the draft is still reviewable"
    assert session.query(Resume).count() == 1
    assert session.query(LlmCall).count() == 0


# --- F2.3.1: an empty extraction is a failure, not a draft ------------------

EMPTY = json.dumps(
    {"language": "es", "headline": "", "summary": "", "roles": [], "credentials": [], "skills": []}
)


def long_cv() -> bytes:
    """A document with enough text that an empty answer is a failure, not a shrug."""
    return docx_bytes(
        "Experiencia",
        "Reduje el cierre mensual un 30% automatizando los reportes con Python y SQL",
        "Mantuve dashboards en producción para tres áreas de negocio distintas",
        "Enseñé SQL a dos analistas junior durante seis meses de trabajo conjunto",
        "Coordiné la migración del reporte manual a un proceso programado",
        "Documenté el proceso y lo dejé en manos del equipo de operaciones",
    )


def test_an_empty_extraction_retries_once_and_reports_the_character_count(
    client, session, career, monkeypatch
):
    empty = Provider(EMPTY)
    monkeypatch.setattr(extract_client, "_post", empty)
    login(client, career.first.email)

    response = upload(client, long_cv())

    page = plain(response.text)
    assert len(empty.calls) == 2, "one retry, then the truth"
    assert re.search(r"\d{3,} caracteres y no volvió nada", page)
    assert "Revisar la extracción" in page, "the findings screen is still the screen"

    rows = session.query(LlmCall).order_by(LlmCall.id).all()
    assert len(rows) == 2
    assert all(row.degenerate is True for row in rows), "neither answer is a cache entry"
    assert all(row.cost_usd > 0 for row in rows), "both were paid for and both stay in the ledger"

    resume = session.query(Resume).one()
    assert resume.draft is None, "archived with no draft to review"
    assert resume.gate is not None, "the findings travel with the file"
    assert resume.file_bytes, "the document is kept for the retry"

    # A second identical upload is not served the empty answer from the cache.
    empty.calls.clear()
    upload(client, long_cv())

    assert len(empty.calls) == 2, "the marked rows are never served back"
    assert session.query(Resume).count() == 2


def test_a_failed_extraction_is_retried_from_the_stored_bytes(client, session, career, monkeypatch):
    """The archive is the point: a retry must not cost another upload."""
    empty = Provider(EMPTY)
    monkeypatch.setattr(extract_client, "_post", empty)
    login(client, career.first.email)
    upload(client, long_cv())
    resume = session.query(Resume).one()

    # The row's own page says what happened and offers the retry, instead of a 404.
    page = client.get(f"/import/{resume.id}").text
    assert "no hay borrador" in plain(page)
    assert f'action="/import/{resume.id}/retry"' in page

    good = Provider()
    monkeypatch.setattr(extract_client, "_post", good)

    retried = client.post(f"/import/{resume.id}/retry", follow_redirects=False)

    assert retried.status_code == 303
    assert retried.headers["location"] == f"/import/{resume.id}"
    session.expire_all()
    assert resume.draft is not None, "the retry wrote the draft the first attempt could not"
    assert "Modelé el riesgo de crédito con Python" in plain(
        client.get(f"/import/{resume.id}").text
    )


# --- F2.5: "nothing changed" is a result, and must be said out loud ----------


def test_a_reimport_of_data_that_already_matches_says_nothing_to_merge(
    client, session, career, provider
):
    """The all-unchanged plan is stated as a result, and the rows stay reachable."""
    login(client, career.first.email)
    data = docx_bytes("Experiencia", "Reduje el cierre un 30% con Python")

    first = upload(client, data)
    client.post(f"{first.headers['location']}/review", data=REVIEWED_FORM)
    client.post(f"{first.headers['location']}/merge", data=MERGE_FORM)

    second = upload(client, data)
    page = client.post(f"{second.headers['location']}/review", data=REVIEWED_FORM)
    text = plain(page.text)

    assert "no hay nada que fusionar" in text
    assert "Se compararon 5 elementos" in text, "new + changed + unchanged, named"
    assert "0 nuevos" in text and "0 cambian" in text, "the counts are still shown"
    assert "<details" in page.text and "open" not in page.text.split("<details")[1][:40], (
        "the candidate list stays below, collapsed"
    )


def test_a_plan_with_no_candidates_says_nothing_was_extracted_to_compare(
    client, session, career, monkeypatch
):
    """The other message: zero candidates points back at the extraction, not at the data."""
    monkeypatch.setattr(extract_client, "_post", Provider(EMPTY))
    login(client, career.first.email)

    # A short document whose answer is empty is a thin draft, not a failed extraction.
    uploaded = upload(client, docx_bytes("Experiencia", "Sin contenido que extraer"))
    page = client.post(f"{uploaded.headers['location']}/review", data={"language": "es"})

    text = plain(page.text)
    assert "No se extrajo nada que comparar" in text
    assert "no hay nada que fusionar" not in text
    assert "Candidatos comparados · 0" in text
    assert "open" in page.text.split("<details")[1][:40], "only the all-unchanged plan collapses"


def test_a_merge_that_writes_nothing_says_the_database_is_unchanged(
    client, session, career, provider
):
    """An all-zero MergeResult is rendered as a statement, not as a blank summary."""
    login(client, career.first.email)
    data = docx_bytes("Experiencia", "Reduje el cierre un 30% con Python")

    # F2.7 made "no track" impossible, so an all-zero result needs a join: a second
    # document under the same track, with nothing new and every slot already filled.
    track_id = str(career.first_track.id)
    first = upload(client, data, track_id=track_id)
    client.post(f"{first.headers['location']}/review", data=REVIEWED_FORM)
    client.post(f"{first.headers['location']}/merge", data=MERGE_FORM)

    second = upload(client, data, track_id=track_id)
    client.post(f"{second.headers['location']}/review", data=REVIEWED_FORM)
    merged = client.post(f"{second.headers['location']}/merge", data={})

    assert merged.status_code == 200
    text = plain(merged.text)
    assert "La fusión no escribió nada" in text
    assert "sin tocar" in text, "the screen says what it left alone because the track had it"


# --- F5.1.2: the invisible-glyph finding, rendered ---------------------------


def glyph_pdf() -> bytes:
    """A clean page with two invisible U+00B7, the operator's own builder artifact."""
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    pdf.setFont("Helvetica", 10)
    pdf.drawString(60, 700, "Experiencia")
    pdf.setFillColorRGB(1, 1, 1)
    pdf.drawString(500, 650, "·")
    pdf.drawString(520, 650, "·")
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def test_the_import_panel_reports_invisible_glyphs_without_capping(
    client, session, career, provider
):
    """A downgrade the user cannot see is indistinguishable from a check that stopped running."""
    login(client, career.first.email)

    response = upload(client, glyph_pdf(), name="cv.pdf", content_type="application/pdf")
    page = client.get(response.headers["location"]) if response.status_code == 303 else response

    text = plain(page.text)
    assert "Glifos invisibles" in text
    assert "2 caracteres invisibles (de puntuación)" in text
    assert "no afectan a cómo un ATS lee este documento" in text
    assert "Bloqueante" not in text, "an export artifact is not a blocker"


# --- 2.2/2.3: the gate reaches the score, and an upload can become a track ----


def test_a_merge_can_turn_the_document_into_a_track(client, session, career, provider):
    """F4 creation path 1: the parsed résumé becomes a track with its bullets selected."""
    login(client, career.first.email)
    data = docx_bytes("Experiencia", "Reduje el cierre un 30% con Python")
    uploaded = upload(client, data)
    client.post(f"{uploaded.headers['location']}/review", data=REVIEWED_FORM)

    created = client.post(
        f"{uploaded.headers['location']}/merge",
        data={**MERGE_FORM, "track_name": "Analista de datos (CV)"},
    )

    assert created.status_code == 200
    assert "Línea de carrera creada: Analista de datos (CV)" in plain(created.text)
    track = session.scalar(select(Track).where(Track.name == "Analista de datos (CV)"))
    assert track is not None
    assert track.lang == "es"
    bullets = [
        bullet.text_es
        for bullet in track_bullets(career.first_profile.id, session, track.id)
    ]
    assert bullets == ["Modelé el riesgo de crédito con Python"], "the document's own bullet"
    session.expire_all()
    assert session.query(Resume).one().track_id == track.id, "the source document is findable"


def test_an_unnamed_merge_still_creates_a_track_named_after_the_file(
    client, session, career, provider
):
    """F2.7: there is no no-track outcome. The tick-box is gone; the name field is
    pre-filled with the file's stem, and an empty one falls back to the same stem."""
    login(client, career.first.email)
    uploaded = upload(client, docx_bytes("Experiencia", "Reduje el cierre un 30% con Python"))
    client.post(f"{uploaded.headers['location']}/review", data=REVIEWED_FORM)

    merged = client.post(f"{uploaded.headers['location']}/merge", data=MERGE_FORM)

    text = plain(merged.text)
    assert "Línea de carrera creada: cv" in text
    assert "Titular: rellenado desde el documento" in text
    assert "Resumen: rellenado desde el documento" in text
    track = session.scalar(select(Track).where(Track.name == "cv"))
    assert track is not None
    session.expire_all()
    assert session.query(Resume).one().track_id == track.id


def test_a_tagged_upload_says_which_track_it_joins_and_how_many_bullets_it_added(
    client, session, career, provider
):
    """F4 path 1's second half: build the track first, then upload a CV under it."""
    login(client, career.first.email)
    uploaded = upload(
        client,
        docx_bytes("Experiencia", "Reduje el cierre un 30% con Python"),
        track_id=str(career.first_track.id),
    )
    assert uploaded.status_code == 303
    review_url = uploaded.headers["location"]

    merge_page = plain(client.post(f"{review_url}/review", data=REVIEWED_FORM).text)
    assert "Los bullets de este documento se unirán a «Analista de datos»" in merge_page

    merged = plain(client.post(f"{review_url}/merge", data=MERGE_FORM).text)
    assert "Bullets añadidos a «Analista de datos»: 1" in merged
    assert "Línea de carrera creada" not in merged, "it joined; it did not create"

    session.expire_all()
    joined = track_bullets(career.first_profile.id, session, career.first_track.id)
    assert [bullet.id for bullet in joined[:1]] == [career.first_bullet.id], "kept its place"
    assert len(joined) == 2, "the document's new bullet was appended"


def test_a_tracks_analysis_is_capped_by_the_document_it_came_from(
    client, session, career, provider
):
    """F5.1: nothing has ever capped a score, because the route passed no gate."""
    login(client, career.first.email)
    # A contact block parked in the header is a blocker: an ATS cannot read the page.
    data = docx_bytes(
        "Experiencia", "Reduje el cierre un 30% con Python", header="persona2@example.com"
    )
    uploaded = upload(client, data)
    client.post(f"{uploaded.headers['location']}/review", data=REVIEWED_FORM)
    client.post(
        f"{uploaded.headers['location']}/merge",
        data={**MERGE_FORM, "track_name": "Analista de datos (CV)"},
    )
    track = session.scalar(select(Track).where(Track.name == "Analista de datos (CV)"))

    posted = client.post(
        "/analyze",
        data={"jd_text": POSTING, "company": "Banco Demo 1", "track_id": str(track.id)},
        follow_redirects=False,
    )

    assert posted.status_code == 303, "the analysis ran"
    application = session.query(Application).one()
    score = application.report["score"]
    assert score["cap"] == 40, "the gate reached the report"
    assert score["cap_reason"] == "contact_in_margin"
    # The limit does not bind here — the w4 sub-scores moved this fixture's total
    # below the cap — which is the case the page must not dress up as a cap.
    assert score["capped"] == score["uncapped"] < score["cap"]
    assert application.scan_readability is not None
    assert [f["check"] for f in application.scan_readability["findings"]] == ["contact_in_margin"]

    page = plain(client.get(posted.headers["location"]).text)
    # Stage 2's rule: the cap sentence fires only when it lowered the score. Here the
    # finding is real and stays visible, but the page must not claim a limit it never
    # applied — this is the «Limitado a 40 … sin límite sería 29» bug class.
    assert "hoy no resta puntos" in page
    assert "Limitado a" not in page
    assert "¿El ATS lee bien tu CV?" in page, "the same findings panel, as the report's question"
    assert "Bloqueante" in page
    assert "No se ejecutaron para este formato" in page, "a DOCX cannot answer every check"


def test_an_analysis_without_a_track_is_not_capped(client, session, career, provider):
    """No track means no source document, which is exactly the behaviour until now."""
    login(client, career.first.email)

    posted = client.post(
        "/analyze", data={"jd_text": POSTING}, follow_redirects=False
    )

    application = session.query(Application).one()
    assert application.report["score"]["cap"] is None
    assert application.scan_readability is None
    assert "El documento limita la puntuación" not in plain(
        client.get(posted.headers["location"]).text
    )
